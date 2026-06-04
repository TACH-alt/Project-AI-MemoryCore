from pathlib import Path
import csv
import sys
import tempfile
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from process_shopee_video import (
    build_ffmpeg_command,
    build_ytdlp_prefix,
    build_ytdlp_command,
    build_caption,
    build_product_folder_name,
    canonicalize_tiktok_url,
    build_tiktok_media_bundle,
    build_tiktok_media_bundle_from_ytdlp_metadata,
    cleanup_non_media_outputs,
    download_with_ytdlp,
    download_tiktok_bundle,
    derive_output_name,
    download_tiktok_photo_audio,
    download_tiktok_thumbnail_with_ytdlp,
    ensure_masterlist_columns,
    extract_tiktok_media_bundle_from_html,
    extension_from_content_type,
    extract_media_urls_from_html,
    group_photo_mode_images,
    find_background_audio,
    filter_media_urls,
    find_executable,
    infer_image_extension,
    is_direct_media_url,
    is_tiktok_video_url,
    load_rows,
    build_photo_mode_status_message,
    normalize_downloaded_image_artifact,
    page_request_headers,
    REQUIRED_COLUMNS,
    replace_path,
    safe_filename,
    select_media_urls_for_download,
    strip_query_for_shopee_product_url,
    tiktok_video_id_from_url,
    update_record,
)


class HelperTests(unittest.TestCase):
    def test_safe_filename_removes_windows_forbidden_characters(self):
        self.assertEqual(safe_filename('A/B:C*D?E"F<G>H|I'), "A-B-C-D-E-F-G-H-I")

    def test_product_folder_uses_product_dash_shop(self):
        self.assertEqual(build_product_folder_name("Mini Fan", "Cool Shop"), "Mini Fan-Cool Shop")

    def test_caption_uses_shop_dash_product(self):
        self.assertEqual(build_caption("Mini Fan", "Cool Shop"), "Cool Shop - Mini Fan")

    def test_output_name_appends_new_before_extension(self):
        self.assertEqual(derive_output_name(Path("demo.mp4")), "demo_new.mp4")

    def test_direct_media_detection(self):
        self.assertTrue(is_direct_media_url("https://example.com/video.mp4?download=1"))
        self.assertTrue(is_direct_media_url("https://example.com/image.jpeg"))
        self.assertFalse(is_direct_media_url("https://shopee.com.my/product/123"))

    def test_extension_from_content_type_handles_audio_types(self):
        self.assertEqual(extension_from_content_type("audio/mp4"), ".m4a")

    def test_load_rows_reads_pending_csv(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as tmp:
            path = Path(tmp) / "master.csv"
            with path.open("w", newline="", encoding="utf-8-sig") as f:
                writer = csv.DictWriter(
                    f,
                    fieldnames=[
                        "Date",
                        "Product Link",
                        "Product Name",
                        "Shop Name",
                        "Source Type",
                        "Status",
                        "Done",
                        "Notes",
                        "Last Processed At",
                        "Output Folder",
                        "Output File",
                    ],
                )
                writer.writeheader()
                writer.writerow(
                    {
                        "Date": "2026-04-23",
                        "Product Link": "https://example.com/video.mp4",
                        "Product Name": "Mini Fan",
                        "Shop Name": "Cool Shop",
                        "Source Type": "Approved Media",
                        "Status": "Pending",
                        "Done": "FALSE",
                    }
                )
            rows = load_rows(path)
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0].row_number, 2)
            self.assertEqual(rows[0].product_name, "Mini Fan")

    def test_find_executable_checks_extra_directories(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as tmp:
            tool = Path(tmp) / "tool.exe"
            tool.write_text("", encoding="utf-8")
            self.assertEqual(find_executable("missing-tool", [Path(tmp) / "tool.exe"]), str(tool))

    def test_ytdlp_command_falls_back_to_python_module(self):
        command = build_ytdlp_command(
            "https://example.com/product",
            Path("raw"),
            ytdlp_path=None,
            python_executable="python",
            allow_path_lookup=False,
        )
        self.assertEqual(command[:3], ["python", "-m", "yt_dlp"])

    def test_ytdlp_command_forces_video_download_for_tiktok(self):
        command = build_ytdlp_command(
            "https://vt.tiktok.com/ZS9hjGD43/",
            Path("raw"),
            ytdlp_path="yt-dlp",
            allow_path_lookup=False,
        )
        self.assertIn("--format", command)
        self.assertIn("bv*+ba/b", command)
        self.assertIn("--merge-output-format", command)
        self.assertIn("mp4", command)

    def test_ytdlp_prefix_falls_back_to_python_module(self):
        self.assertEqual(
            build_ytdlp_prefix(ytdlp_path=None, python_executable="python", allow_path_lookup=False),
            ["python", "-m", "yt_dlp"],
        )

    def test_canonicalize_tiktok_url_extracts_long_video_url_from_oembed(self):
        class FakeResponse:
            def __init__(self, body: bytes):
                self._body = body

            def read(self) -> bytes:
                return self._body

            def __enter__(self):
                return self

            def __exit__(self, exc_type, exc, tb):
                return False

        def fake_request(url: str, headers=None):
            self.assertIn("oembed", url)
            return FakeResponse(
                b'{"author_url":"https://www.tiktok.com/@shopdemo","html":"<blockquote cite=\\"https://www.tiktok.com/@shopdemo/video/7490123456789012345\\"></blockquote>"}'
            )

        result = canonicalize_tiktok_url(
            "https://vt.tiktok.com/ZS9hjGD43/",
            request_fn=fake_request,
        )

        self.assertEqual(
            result,
            "https://www.tiktok.com/@shopdemo/video/7490123456789012345",
        )

    def test_tiktok_video_id_from_url_extracts_numeric_id(self):
        self.assertEqual(
            tiktok_video_id_from_url("https://www.tiktok.com/@shopdemo/video/7490123456789012345"),
            "7490123456789012345",
        )

    def test_is_tiktok_video_url_accepts_photo_urls(self):
        self.assertTrue(
            is_tiktok_video_url("https://www.tiktok.com/@shopdemo/photo/7490123456789012345")
        )

    def test_tiktok_video_id_from_url_extracts_numeric_id_from_photo_url(self):
        self.assertEqual(
            tiktok_video_id_from_url("https://www.tiktok.com/@shopdemo/photo/7490123456789012345"),
            "7490123456789012345",
        )

    def test_cleanup_non_media_outputs_removes_audio_only_tiktok_artifacts(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as tmp:
            raw_dir = Path(tmp) / "raw"
            raw_dir.mkdir()
            keep = raw_dir / "video.mp4"
            remove = raw_dir / "audio-only.m4a"
            keep.write_bytes(b"video")
            remove.write_bytes(b"audio")

            removed = cleanup_non_media_outputs(raw_dir)

            self.assertEqual(removed, [remove])
            self.assertTrue(keep.exists())
            self.assertFalse(remove.exists())

    def test_cleanup_non_media_outputs_keeps_background_audio(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as tmp:
            raw_dir = Path(tmp) / "raw"
            raw_dir.mkdir()
            keep = raw_dir / "background.m4a"
            keep.write_bytes(b"audio")

            removed = cleanup_non_media_outputs(raw_dir)

            self.assertEqual(removed, [])
            self.assertTrue(keep.exists())

    def test_infer_image_extension_detects_jpeg_bytes(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as tmp:
            artifact = Path(tmp) / "thumb.image"
            artifact.write_bytes(b"\xff\xd8\xff\xe0" + b"\x00" * 8)
            self.assertEqual(infer_image_extension(artifact), ".jpg")

    def test_normalize_downloaded_image_artifact_renames_image_suffix(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as tmp:
            artifact = Path(tmp) / "thumb.image"
            artifact.write_bytes(b"\xff\xd8\xff\xe0" + b"\x00" * 8)
            normalized = normalize_downloaded_image_artifact(artifact)
            self.assertEqual(normalized.suffix.lower(), ".jpg")
            self.assertTrue(normalized.exists())
            self.assertFalse(artifact.exists())

    def test_normalize_downloaded_image_artifact_replaces_existing_target(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as tmp:
            artifact = Path(tmp) / "thumb.image"
            existing = Path(tmp) / "thumb.jpg"
            existing.write_bytes(b"old")
            artifact.write_bytes(b"\xff\xd8\xff\xe0" + b"new-bytes")

            normalized = normalize_downloaded_image_artifact(artifact)

            self.assertEqual(normalized, existing)
            self.assertEqual(existing.read_bytes(), b"\xff\xd8\xff\xe0new-bytes")
            self.assertFalse(artifact.exists())

    def test_download_tiktok_thumbnail_with_ytdlp_returns_normalized_image(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as tmp:
            raw_dir = Path(tmp) / "raw"
            raw_dir.mkdir()

            def fake_run(command, capture_output, text, check):
                self.assertIn("--write-thumbnail", command)
                self.assertIn("--force-overwrites", command)
                self.assertIn(str(raw_dir), command)
                (raw_dir / "7129084799339351323.image").write_bytes(b"\xff\xd8\xff\xe0" + b"\x00" * 8)

                class Result:
                    returncode = 0
                return Result()

            from unittest.mock import patch

            with patch("process_shopee_video.subprocess.run", side_effect=fake_run):
                images = download_tiktok_thumbnail_with_ytdlp(
                    "https://www.tiktok.com/@shopdemo/video/7490123456789012345",
                    raw_dir,
                )

            self.assertEqual(len(images), 1)
            self.assertEqual(images[0].suffix.lower(), ".jpg")
            self.assertTrue(images[0].exists())

    def test_download_tiktok_thumbnail_with_ytdlp_reuses_overwritten_existing_image(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as tmp:
            raw_dir = Path(tmp) / "raw"
            raw_dir.mkdir()
            existing = raw_dir / "7490123456789012345.jpg"
            existing.write_bytes(b"old")

            def fake_run(command, capture_output, text, check):
                existing.write_bytes(b"\xff\xd8\xff\xe0" + b"updated")

                class Result:
                    returncode = 0
                return Result()

            from unittest.mock import patch

            with patch("process_shopee_video.subprocess.run", side_effect=fake_run):
                images = download_tiktok_thumbnail_with_ytdlp(
                    "https://www.tiktok.com/@shopdemo/video/7490123456789012345",
                    raw_dir,
                )

            self.assertEqual(images, [existing])
            self.assertEqual(existing.read_bytes(), b"\xff\xd8\xff\xe0updated")

    def test_download_tiktok_photo_audio_returns_downloaded_native_audio(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as tmp:
            raw_dir = Path(tmp) / "raw"
            raw_dir.mkdir()

            def fake_run(command, capture_output, text, check):
                self.assertIn("--format", command)
                self.assertIn("bestaudio", command)
                self.assertIn("--force-overwrites", command)
                self.assertNotIn("-x", command)
                (raw_dir / "background.m4a").write_bytes(b"audio")

                class Result:
                    returncode = 0
                return Result()

            from unittest.mock import patch

            with patch("process_shopee_video.subprocess.run", side_effect=fake_run):
                audio = download_tiktok_photo_audio(
                    "https://www.tiktok.com/@shopdemo/video/7490123456789012345",
                    raw_dir,
                )

            self.assertIsNotNone(audio)
            self.assertEqual(audio.name, "background.m4a")
            self.assertTrue(audio.exists())

    def test_extract_tiktok_media_bundle_from_html_prefers_video_urls(self):
        body = """
        <script id="__UNIVERSAL_DATA_FOR_REHYDRATION__" type="application/json">
        {"__DEFAULT_SCOPE__":{"webapp.video-detail":{"itemInfo":{"itemStruct":{
            "id":"7490123456789012345",
            "video":{"playAddr":"https://cdn.example.com/play.mp4","downloadAddr":"https://cdn.example.com/download.mp4"},
            "music":{"playUrl":"https://cdn.example.com/music.m4a"}
        }}}}}
        </script>
        """

        bundle = extract_tiktok_media_bundle_from_html(
            "https://www.tiktok.com/@shopdemo/video/7490123456789012345",
            body,
        )

        self.assertIsNotNone(bundle)
        self.assertEqual(bundle["kind"], "video")
        self.assertEqual(bundle["post_type"], "video")
        self.assertEqual(bundle["video_id"], "7490123456789012345")
        self.assertEqual(
            bundle["video_urls"],
            ["https://cdn.example.com/download.mp4", "https://cdn.example.com/play.mp4"],
        )
        self.assertEqual(bundle["audio_url"], "https://cdn.example.com/music.m4a")

    def test_extract_tiktok_media_bundle_from_html_handles_photo_mode(self):
        body = """
        <script id="__UNIVERSAL_DATA_FOR_REHYDRATION__" type="application/json">
        {"__DEFAULT_SCOPE__":{"webapp.video-detail":{"itemInfo":{"itemStruct":{
            "id":"7490123456789012345",
            "video":{"playAddr":"","downloadAddr":"","originCover":"https://cdn.example.com/origin.image",
                     "zoomCover":{"240":"https://cdn.example.com/240.jpeg","960":"https://cdn.example.com/960.jpeg"}},
            "music":{"playUrl":"https://cdn.example.com/music.m4a"}
        }}}}}
        </script>
        """

        bundle = extract_tiktok_media_bundle_from_html(
            "https://www.tiktok.com/@shopdemo/video/7490123456789012345",
            body,
        )

        self.assertIsNotNone(bundle)
        self.assertEqual(bundle["kind"], "photo_mode")
        self.assertEqual(bundle["post_type"], "photo_mode")
        self.assertEqual(bundle["video_id"], "7490123456789012345")
        self.assertEqual(
            bundle["image_urls"],
            ["https://cdn.example.com/origin.image"],
        )
        self.assertEqual(bundle["audio_url"], "https://cdn.example.com/music.m4a")

    def test_extract_tiktok_media_bundle_from_html_marks_unknown_when_no_media_assets_exist(self):
        body = """
        <script id="__UNIVERSAL_DATA_FOR_REHYDRATION__" type="application/json">
        {"__DEFAULT_SCOPE__":{"webapp.video-detail":{"itemInfo":{"itemStruct":{
            "id":"7490123456789012345",
            "video":{"playAddr":"","downloadAddr":""},
            "music":{"playUrl":""}
        }}}}}
        </script>
        """

        bundle = extract_tiktok_media_bundle_from_html(
            "https://www.tiktok.com/@shopdemo/photo/7490123456789012345",
            body,
        )

        self.assertIsNotNone(bundle)
        self.assertEqual(bundle["kind"], "unknown")
        self.assertEqual(bundle["post_type"], "unknown")
        self.assertEqual(bundle["video_id"], "7490123456789012345")
        self.assertEqual(bundle["video_urls"], [])
        self.assertEqual(bundle["image_urls"], [])
        self.assertIsNone(bundle["audio_url"])

    def test_extract_tiktok_media_bundle_from_html_supports_next_data_payload(self):
        body = """
        <script id="__NEXT_DATA__" type="application/json">
        {"props":{"pageProps":{"itemInfo":{"itemStruct":{
            "id":"7490123456789012345",
            "video":{"playAddr":"https://cdn.example.com/play.mp4","downloadAddr":"https://cdn.example.com/download.mp4"},
            "music":{"playUrl":"https://cdn.example.com/music.m4a"}
        }}}}}
        </script>
        """

        bundle = extract_tiktok_media_bundle_from_html(
            "https://www.tiktok.com/@shopdemo/video/7490123456789012345",
            body,
        )

        self.assertIsNotNone(bundle)
        self.assertEqual(bundle["kind"], "video")
        self.assertEqual(
            bundle["video_urls"],
            ["https://cdn.example.com/download.mp4", "https://cdn.example.com/play.mp4"],
        )

    def test_extract_tiktok_media_bundle_from_html_supports_sigi_state_payload(self):
        body = """
        <script id="SIGI_STATE" type="application/json">
        {"ItemModule":{"7490123456789012345":{
            "id":"7490123456789012345",
            "video":{"playAddr":"https://cdn.example.com/play.mp4","downloadAddr":"https://cdn.example.com/download.mp4"},
            "music":{"playUrl":"https://cdn.example.com/music.m4a"}
        }}}
        </script>
        """

        bundle = extract_tiktok_media_bundle_from_html(
            "https://www.tiktok.com/@shopdemo/video/7490123456789012345",
            body,
        )

        self.assertIsNotNone(bundle)
        self.assertEqual(bundle["kind"], "video")
        self.assertEqual(bundle["video_id"], "7490123456789012345")

    def test_build_tiktok_media_bundle_supports_node_share_payload(self):
        payload = {
            "statusCode": 0,
            "itemInfo": {
                "itemStruct": {
                    "id": "7490123456789012345",
                    "video": {
                        "playAddr": "https://cdn.example.com/play.mp4",
                        "downloadAddr": "https://cdn.example.com/download.mp4",
                    },
                    "music": {"playUrl": "https://cdn.example.com/music.m4a"},
                }
            },
        }

        bundle = build_tiktok_media_bundle(
            payload,
            "https://www.tiktok.com/@shopdemo/video/7490123456789012345",
        )

        self.assertIsNotNone(bundle)
        self.assertEqual(bundle["kind"], "video")
        self.assertEqual(
            bundle["video_urls"],
            ["https://cdn.example.com/download.mp4", "https://cdn.example.com/play.mp4"],
        )

    def test_build_tiktok_media_bundle_from_ytdlp_metadata_prefers_video_formats(self):
        metadata = {
            "id": "7490123456789012345",
            "formats": [
                {
                    "format_id": "audio",
                    "url": "https://cdn.example.com/audio.m4a",
                    "vcodec": "none",
                    "acodec": "aac",
                    "ext": "m4a",
                },
                {
                    "format_id": "video",
                    "url": "https://cdn.example.com/video.mp4",
                    "vcodec": "h264",
                    "acodec": "aac",
                    "ext": "mp4",
                    "height": 1080,
                },
            ],
            "thumbnail": "https://cdn.example.com/thumb.jpg",
        }

        bundle = build_tiktok_media_bundle_from_ytdlp_metadata(
            metadata,
            "https://www.tiktok.com/@shopdemo/video/7490123456789012345",
        )

        self.assertIsNotNone(bundle)
        self.assertEqual(bundle["kind"], "video")
        self.assertEqual(bundle["video_urls"], ["https://cdn.example.com/video.mp4"])
        self.assertEqual(bundle["image_urls"], ["https://cdn.example.com/thumb.jpg"])

    def test_build_tiktok_media_bundle_from_ytdlp_metadata_marks_photo_mode_for_audio_and_thumbnails(self):
        metadata = {
            "id": "7490123456789012345",
            "formats": [
                {
                    "format_id": "audio",
                    "url": "https://cdn.example.com/audio.m4a",
                    "vcodec": "none",
                    "acodec": "aac",
                    "ext": "m4a",
                }
            ],
            "thumbnails": [
                {"id": "cover", "url": "https://cdn.example.com/cover.jpg", "preference": -1},
                {"id": "originCover", "url": "https://cdn.example.com/origin.jpg", "preference": 1},
            ],
        }

        bundle = build_tiktok_media_bundle_from_ytdlp_metadata(
            metadata,
            "https://www.tiktok.com/@shopdemo/video/7490123456789012345",
        )

        self.assertIsNotNone(bundle)
        self.assertEqual(bundle["kind"], "photo_mode")
        self.assertEqual(bundle["video_urls"], [])
        self.assertEqual(bundle["image_urls"], ["https://cdn.example.com/origin.jpg", "https://cdn.example.com/cover.jpg"])
        self.assertEqual(bundle["audio_url"], "https://cdn.example.com/audio.m4a")

    def test_download_tiktok_bundle_downloads_photo_mode_images_and_audio(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as tmp:
            raw_dir = Path(tmp) / "raw"
            raw_dir.mkdir()
            bundle = {
                "kind": "photo_mode",
                "video_id": "7490123456789012345",
                "image_urls": [
                    "https://cdn.example.com/one.jpeg",
                    "https://cdn.example.com/two.jpeg",
                ],
                "audio_url": "https://cdn.example.com/music.m4a",
            }

            created: list[Path] = []

            def fake_download(url, target_dir, fallback_name="media"):
                suffix = ".m4a" if "music" in url else ".jpg"
                target = target_dir / f"{fallback_name}{suffix}"
                target.write_bytes(b"demo")
                created.append(target)
                return target

            from unittest.mock import patch

            with patch("process_shopee_video.download_url", side_effect=fake_download):
                media = download_tiktok_bundle(bundle, raw_dir)

            self.assertEqual(
                [path.name for path in media],
                ["7490123456789012345.jpg", "7490123456789012345-2.jpg"],
            )
            self.assertTrue(all(path.exists() for path in media))
            self.assertTrue((raw_dir / "background.m4a").exists())

    def test_download_tiktok_bundle_prefers_direct_video_target(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as tmp:
            raw_dir = Path(tmp) / "raw"
            raw_dir.mkdir()
            bundle = {
                "post_type": "video",
                "video_id": "7490123456789012345",
                "video_urls": ["https://cdn.example.com/original-name.mp4"],
                "image_urls": [],
                "audio_url": "",
            }

            def fake_download(url, target_dir, fallback_name="media"):
                target = target_dir / "original-name.mp4"
                target.write_bytes(b"demo")
                return target

            from unittest.mock import patch

            with patch("process_shopee_video.download_url", side_effect=fake_download):
                media = download_tiktok_bundle(bundle, raw_dir)

            self.assertEqual([path.name for path in media], ["7490123456789012345.mp4"])
            self.assertTrue((raw_dir / "7490123456789012345.mp4").exists())

    def test_download_tiktok_bundle_keeps_multiple_photo_assets_in_order(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as tmp:
            raw_dir = Path(tmp) / "raw"
            raw_dir.mkdir()
            bundle = {
                "post_type": "photo_mode",
                "video_id": "7490123456789012345",
                "video_urls": [],
                "image_urls": [
                    "https://cdn.example.com/one.jpeg",
                    "https://cdn.example.com/two.jpeg",
                    "https://cdn.example.com/three.jpeg",
                ],
                "audio_url": "https://cdn.example.com/music.m4a",
            }

            def fake_download(url, target_dir, fallback_name="media"):
                suffix = ".m4a" if "music" in url else ".jpg"
                target = target_dir / f"{fallback_name}{suffix}"
                target.write_bytes(b"demo")
                return target

            from unittest.mock import patch

            with patch("process_shopee_video.download_url", side_effect=fake_download):
                media = download_tiktok_bundle(bundle, raw_dir)

            self.assertEqual(
                [path.name for path in media],
                ["7490123456789012345.jpg", "7490123456789012345-2.jpg", "7490123456789012345-3.jpg"],
            )
            self.assertTrue((raw_dir / "background.m4a").exists())

    def test_download_with_ytdlp_replaces_tiktok_thumbnail_result_with_structured_bundle(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as tmp:
            raw_dir = Path(tmp) / "raw"
            raw_dir.mkdir()

            def fake_run(command, capture_output, text, check):
                (raw_dir / "thumb.jpg").write_bytes(b"thumb")

                class Result:
                    returncode = 0
                return Result()

            def fake_bundle(url):
                return {
                    "kind": "photo_mode",
                    "video_id": "7490123456789012345",
                    "image_urls": ["https://cdn.example.com/one.jpeg"],
                    "audio_url": "https://cdn.example.com/music.m4a",
                }

            def fake_download_bundle(bundle, directory):
                target = directory / "7490123456789012345-1.jpg"
                target.write_bytes(b"better")
                return [target]

            from unittest.mock import patch

            with (
                patch("process_shopee_video.subprocess.run", side_effect=fake_run),
                patch("process_shopee_video.discover_tiktok_media_bundle", side_effect=fake_bundle),
                patch("process_shopee_video.download_tiktok_bundle", side_effect=fake_download_bundle),
            ):
                media = download_with_ytdlp(
                    "https://www.tiktok.com/@shopdemo/video/7490123456789012345",
                    raw_dir,
                )

            self.assertEqual([path.name for path in media], ["7490123456789012345-1.jpg"])
            self.assertFalse((raw_dir / "thumb.jpg").exists())

    def test_group_photo_mode_sequence_collects_ordered_images(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as tmp:
            raw_dir = Path(tmp) / "raw"
            raw_dir.mkdir()
            for name in ["7490.jpg", "7490-2.jpg", "7490-3.jpg", "other.jpg"]:
                (raw_dir / name).write_bytes(b"img")
            files = group_photo_mode_images(raw_dir / "7490.jpg", raw_dir)
            self.assertEqual([path.name for path in files], ["7490.jpg", "7490-2.jpg", "7490-3.jpg"])

    def test_photo_mode_status_message_mentions_slideshow(self):
        message = build_photo_mode_status_message(image_count=3, edited_count=1, has_audio=True)
        self.assertEqual(
            message,
            "Downloaded 4 media file(s); edited 1 slideshow video file(s) (photo-mode post with background audio)",
        )

    def test_find_background_audio_prefers_existing_supported_file(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as tmp:
            raw_dir = Path(tmp) / "raw"
            raw_dir.mkdir()
            audio = raw_dir / "background.m4a"
            audio.write_bytes(b"audio")

            self.assertEqual(find_background_audio(raw_dir), audio)

    def test_replace_path_overwrites_existing_file(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as tmp:
            base = Path(tmp)
            target = base / "clip.mp4"
            target.write_bytes(b"old")

            resolved = replace_path(base, "clip.mp4")
            resolved.write_bytes(b"new")

            self.assertEqual(resolved, target)
            self.assertEqual(target.read_bytes(), b"new")

    def test_filter_media_urls_skips_site_chrome_assets(self):
        urls = filter_media_urls(
            [
                "https://example.com/icon_favicon_1_32.png",
                "https://example.com/ios_icon_180x180.png",
                "https://example.com/ios_splash_screen_1080x2340.png",
                "https://deo.shopeemobile.com/shopee/shopee-mobilemall-live-sg/productdetailspage/share.png",
                "https://example.com/logo.png",
                "https://example.com/product-demo.mp4",
            ]
        )
        self.assertEqual(urls, ["https://example.com/product-demo.mp4"])

    def test_shopee_product_url_query_is_removed_before_scraping(self):
        cleaned = strip_query_for_shopee_product_url(
            "https://shopee.com.my/Electrova-Cordless-UV-Mite-Remover-i.297985880.25153849822"
            "?extraParams=%7B%22display_model_id%22%3A250102961680%7D"
        )
        self.assertEqual(
            cleaned,
            "https://shopee.com.my/Electrova-Cordless-UV-Mite-Remover-i.297985880.25153849822",
        )

    def test_page_request_headers_use_mobile_browser_for_shopee_pages(self):
        headers = page_request_headers("https://shopee.com.my/product/297985880/25153849822")
        self.assertIn("iPhone", headers["User-Agent"])
        self.assertIn("en-MY", headers["Accept-Language"])

    def test_extract_media_urls_from_shopee_seo_html_prefers_product_media(self):
        html = """
        <meta property="og:image" content="https://down-my.img.susercontent.com/file/my-11134207-demo@resize_w1200_nl.webp">
        <script>{"video_info_list":[{"formats":[{"url":"https://down-ws-sg.vod.susercontent.com/api/v4/11110105/mms/demo.1600325.mp4"}]}]}</script>
        <img src="https://deo.shopeemobile.com/shopee/shopee-mobilemall-live-sg/assets/icon_favicon_1_32.png">
        <style>.share{background-image:url(https://deo.shopeemobile.com/shopee/shopee-mobilemall-live-sg/productdetailspage/share.png)}</style>
        """
        urls = extract_media_urls_from_html("https://shopee.com.my/product/297985880/25153849822", html)
        self.assertEqual(
            urls,
            [
                "https://down-ws-sg.vod.susercontent.com/api/v4/11110105/mms/demo.1600325.mp4",
                "https://down-my.img.susercontent.com/file/my-11134207-demo@resize_w1200_nl.webp",
            ],
        )

    def test_select_media_urls_downloads_single_best_video_when_available(self):
        urls = select_media_urls_for_download(
            [
                "https://down-ws-sg.vod.susercontent.com/api/v4/11110105/mms/demo.1600318.mp4",
                "https://down-ws-sg.vod.susercontent.com/api/v4/11110105/mms/demo.1600325.mp4",
                "https://down-my.img.susercontent.com/file/my-11134207-demo@resize_w1200_nl.webp",
            ]
        )
        self.assertEqual(
            urls,
            ["https://down-ws-sg.vod.susercontent.com/api/v4/11110105/mms/demo.1600325.mp4"],
        )

    def test_ffmpeg_command_uses_complex_filter_for_composited_canvas(self):
        command = build_ffmpeg_command(
            "ffmpeg",
            Path("raw.png"),
            Path("edited_new.mp4"),
            "Shop - Product",
        )
        self.assertIn("-filter_complex", command)
        self.assertNotIn("-vf", command)

    def test_update_record_clears_stale_output_file_when_no_outputs(self):
        record = {
            "Status": "Edited",
            "Done": "TRUE",
            "Notes": "old",
            "Last Processed At": "",
            "Output Folder": "old-folder",
            "Output File": "old-file.mp4",
        }
        update_record(record, "Failed", "No downloadable media found", Path("folder"), [])
        self.assertEqual(record["Status"], "Failed")
        self.assertEqual(record["Done"], "FALSE")
        self.assertEqual(record["Output File"], "")

    def test_update_record_marks_done_when_downloaded_or_edited(self):
        for status in ("Downloaded", "Edited"):
            record = {
                "Status": "Pending",
                "Done": "FALSE",
                "Notes": "",
                "Last Processed At": "",
                "Output Folder": "",
                "Output File": "",
            }
            update_record(record, status, "ok", Path("folder"), [Path("video_new.mp4")])
            self.assertEqual(record["Done"], "TRUE")

    def test_ensure_masterlist_columns_adds_done_to_existing_csv(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as tmp:
            path = Path(tmp) / "master.csv"
            path.write_text(
                "Date,Product Link,Product Name,Shop Name,Source Type,Status,Notes,Last Processed At,Output Folder,Output File\n"
                "2026-04-23,https://example.com/video.mp4,Mini Fan,Cool Shop,Approved Media,Pending,,,,\n",
                encoding="utf-8-sig",
            )
            ensure_masterlist_columns(path)
            with path.open("r", newline="", encoding="utf-8-sig") as handle:
                rows = list(csv.DictReader(handle))
            self.assertIn("Done", rows[0])
            self.assertEqual(rows[0]["Done"], "FALSE")

    def test_write_masterlist_syncs_existing_workbook(self):
        from openpyxl import Workbook, load_workbook
        from process_shopee_video import write_masterlist

        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as tmp:
            csv_path = Path(tmp) / "master.csv"
            workbook_path = Path(tmp) / "master.xlsx"
            workbook = Workbook()
            workbook.active.title = "List of Link Shopee VIdeo"
            workbook.save(workbook_path)

            write_masterlist(csv_path, [{
                "Date": "2026-04-23",
                "Product Link": "https://example.com/video.mp4",
                "Product Name": "Mini Fan",
                "Shop Name": "Cool Shop",
                "Source Type": "Approved Media",
                "Status": "Downloaded",
                "Done": "TRUE",
                "Notes": "ok",
                "Last Processed At": "",
                "Output Folder": "",
                "Output File": "",
            }])

            updated = load_workbook(workbook_path)
            sheet = updated["List of Link Shopee VIdeo"]
            headers = [sheet.cell(1, column).value for column in range(1, sheet.max_column + 1)]
            self.assertEqual(headers, REQUIRED_COLUMNS)
            self.assertEqual(sheet.cell(2, 1).value, "ASV-0000001")
            self.assertEqual(sheet.cell(2, 8).value, "TRUE")

    def test_no_media_message_is_actionable_for_shopee_pages(self):
        from process_shopee_video import no_media_message

        self.assertEqual(
            no_media_message("https://shopee.com.my/product/32632559/41452031782"),
            "Shopee page did not expose downloadable product media; use an approved seller media/direct file link",
        )


class BulkIntakeTests(unittest.TestCase):
    def test_parse_url_lines_strips_comments_prefixes_and_dupes(self):
        from process_shopee_video import parse_url_lines

        text = """
# header comment
1. https://shopee.com.my/Foo-i.1.2
- https://www.tiktok.com/@u/video/123

https://shopee.com.my/Foo-i.1.2
"""
        self.assertEqual(
            parse_url_lines(text),
            ["https://shopee.com.my/Foo-i.1.2", "https://www.tiktok.com/@u/video/123"],
        )

    def test_shopee_slug_extracts_product_name(self):
        from process_shopee_video import shopee_slug_to_product_name

        self.assertEqual(
            shopee_slug_to_product_name("https://shopee.com.my/Cordless-UV-Mite-Vacuum-i.123.456"),
            "Cordless UV Mite Vacuum",
        )
        self.assertEqual(shopee_slug_to_product_name("https://shopee.com.my/no-marker"), "")

    def test_is_shopee_shortlink(self):
        from process_shopee_video import is_shopee_shortlink

        self.assertTrue(is_shopee_shortlink("https://s.shopee.com.my/abc"))
        self.assertFalse(is_shopee_shortlink("https://shopee.com.my/x-i.1.2"))

    def test_derive_metadata_rejects_shortlink(self):
        from process_shopee_video import derive_metadata_from_url

        meta = derive_metadata_from_url("https://s.shopee.com.my/abc")
        self.assertIn("short links", meta["error"])
        self.assertEqual(meta["product_name"], "")

    def test_derive_metadata_tiktok_uses_username_and_video_id(self):
        from process_shopee_video import derive_metadata_from_url

        meta = derive_metadata_from_url("https://www.tiktok.com/@bargain/video/7350000000000000000")
        self.assertEqual(meta["platform"], "TikTok")
        self.assertEqual(meta["shop_name"], "@bargain")
        self.assertIn("7350000000000000000", meta["product_name"])

    def test_bulk_intake_appends_pending_rows_and_dedupes(self):
        from process_shopee_video import bulk_intake_rows, read_masterlist

        with tempfile.TemporaryDirectory() as tmp:
            csv_path = Path(tmp) / "masterlist.csv"
            results = bulk_intake_rows(
                [
                    "https://shopee.com.my/Foo-Bar-i.1.2",
                    "https://www.tiktok.com/@u/video/9999999999999999999",
                    "https://shopee.com.my/Foo-Bar-i.1.2",  # duplicate within batch
                    "https://s.shopee.com.my/zzz",  # rejected
                ],
                csv_path,
                today="2026-05-05",
            )
            statuses = [r["status"] for r in results]
            self.assertEqual(statuses.count("added"), 2)
            self.assertEqual(statuses.count("skipped"), 1)
            self.assertEqual(statuses.count("error"), 1)

            rows = read_masterlist(csv_path)
            self.assertEqual(len(rows), 2)
            self.assertEqual(rows[0]["Status"], "Pending")
            self.assertEqual(rows[0]["Date"], "2026-05-05")
            self.assertTrue(rows[0]["Record ID"].startswith("ASV-"))
            self.assertTrue(rows[1]["Record ID"].startswith("ATV-"))

            # Calling again with same URLs is a no-op (skipped)
            results2 = bulk_intake_rows(
                ["https://shopee.com.my/Foo-Bar-i.1.2"], csv_path, today="2026-05-05"
            )
            self.assertEqual(results2[0]["status"], "skipped")


class RenderModeTests(unittest.TestCase):
    def test_is_vertical_9_16_accepts_native_resolutions(self):
        from process_shopee_video import is_vertical_9_16

        self.assertTrue(is_vertical_9_16(1080, 1920))
        self.assertTrue(is_vertical_9_16(720, 1280))
        # Square — not vertical
        self.assertFalse(is_vertical_9_16(720, 720))
        # Landscape — not vertical
        self.assertFalse(is_vertical_9_16(1920, 1080))
        # Too small
        self.assertFalse(is_vertical_9_16(360, 640))

    def test_passthrough_filter_skips_blur_and_pad(self):
        command = build_ffmpeg_command(
            "ffmpeg",
            Path("input.mp4"),
            Path("output.mp4"),
            "Caption",
            render_mode="passthrough",
        )
        graph = command[command.index("-filter_complex") + 1]
        self.assertNotIn("boxblur", graph)
        self.assertNotIn("pad=", graph)
        self.assertIn("scale=1080:1920", graph)
        self.assertIn("drawtext", graph)
        # Audio is preserved via the optional source map
        self.assertIn("0:a?", command)


class CaptionPackTests(unittest.TestCase):
    def test_build_caption_pack_prompt_includes_required_fields(self):
        from process_shopee_video import build_caption_pack_prompt

        record = {
            "Product Name": "Cordless UV Mite Vacuum",
            "Shop Name": "Electrova",
            "Product Link": "https://shopee.com.my/x-i.1.2",
            "Notes": "approved seller media",
        }
        prompt = build_caption_pack_prompt(record, "Shopee")
        for needle in ("Cordless UV Mite Vacuum", "Electrova", "Shopee", "tiktok", "reels", "shorts_title", "hashtags", "JSON"):
            self.assertIn(needle, prompt, f"prompt missing '{needle}'")

    def test_extract_json_object_handles_markdown_fences(self):
        from process_shopee_video import extract_json_object

        wrapped = "```json\n{\"tiktok\": \"hi\", \"hashtags\": [\"#a\", \"#b\"]}\n```"
        self.assertEqual(extract_json_object(wrapped), {"tiktok": "hi", "hashtags": ["#a", "#b"]})

    def test_extract_json_object_handles_trailing_prose(self):
        from process_shopee_video import extract_json_object

        noisy = "Sure! Here is the JSON:\n{\"tiktok\":\"hi\"}\nLet me know if you need more."
        self.assertEqual(extract_json_object(noisy), {"tiktok": "hi"})

    def test_extract_json_object_returns_none_on_garbage(self):
        from process_shopee_video import extract_json_object

        self.assertIsNone(extract_json_object("no json here"))
        self.assertIsNone(extract_json_object(""))

    def test_format_captions_pack_renders_all_sections(self):
        from process_shopee_video import format_captions_pack

        pack = {
            "tiktok": "Hook line + CTA",
            "reels": "Different copy",
            "shorts_caption": "Watch til end",
            "shorts_title": "Best deal",
            "shopee": "Long persuasive copy",
            "hashtags": ["#shopeefinds", "#tiktokmademebuyit"],
        }
        record = {"Record ID": "ASV-0000001", "Product Name": "Foo", "Shop Name": "Bar", "Product Link": "https://x"}
        text = format_captions_pack(pack, record)
        for needle in ("ASV-0000001", "TikTok", "Reels", "Shorts (YouTube)", "Shorts Title", "Shopee Affiliate", "Hashtags", "#shopeefinds"):
            self.assertIn(needle, text)

    def test_call_llm_falls_back_on_401(self):
        from process_shopee_video import call_llm

        calls = []

        def fake_transport(url, payload, timeout):
            calls.append(payload["model"])
            if payload["model"] == "primary":
                return 401, {"error": {"message": "auth expired"}}
            return 200, {"choices": [{"message": {"content": "fallback worked"}}]}

        result = call_llm("hi", model="primary", fallback_model="backup", transport=fake_transport)
        self.assertEqual(result, "fallback worked")
        self.assertEqual(calls, ["primary", "backup"])

    def test_call_llm_raises_on_double_failure(self):
        from process_shopee_video import call_llm

        def fake_transport(url, payload, timeout):
            return 500, {"error": {"message": "router down"}}

        with self.assertRaises(RuntimeError):
            call_llm("hi", model="a", fallback_model="b", transport=fake_transport)

    def test_write_captions_pack_writes_formatted_file_on_valid_json(self):
        from process_shopee_video import write_captions_pack

        with tempfile.TemporaryDirectory() as tmp:
            product_dir = Path(tmp) / "p"

            def fake_transport(url, payload, timeout):
                return 200, {"choices": [{"message": {"content": '{"tiktok":"a","reels":"b","shorts_caption":"c","shorts_title":"d","shopee":"e","hashtags":["#x"]}'}}]}

            path, err = write_captions_pack(
                product_dir,
                {"Record ID": "ASV-1", "Product Name": "Foo", "Shop Name": "Bar", "Product Link": "https://x"},
                "Shopee",
                transport=fake_transport,
            )
            self.assertIsNone(err)
            self.assertIsNotNone(path)
            self.assertTrue(path.exists())
            text = path.read_text(encoding="utf-8")
            self.assertIn("=== TikTok ===", text)
            self.assertIn("#x", text)

    def test_write_captions_pack_saves_raw_when_json_invalid(self):
        from process_shopee_video import write_captions_pack

        with tempfile.TemporaryDirectory() as tmp:
            product_dir = Path(tmp) / "p"

            def fake_transport(url, payload, timeout):
                return 200, {"choices": [{"message": {"content": "the model rambled and forgot to JSON"}}]}

            path, err = write_captions_pack(
                product_dir,
                {"Record ID": "ASV-1", "Product Name": "Foo", "Shop Name": "Bar", "Product Link": "https://x"},
                "Shopee",
                transport=fake_transport,
            )
            self.assertIsNotNone(err)
            self.assertIsNotNone(path)
            self.assertIn("WARNING", path.read_text(encoding="utf-8"))

    def test_write_captions_pack_returns_error_on_llm_failure(self):
        from process_shopee_video import write_captions_pack

        with tempfile.TemporaryDirectory() as tmp:
            product_dir = Path(tmp) / "p"

            def fake_transport(url, payload, timeout):
                return 500, {"error": {"message": "down"}}

            path, err = write_captions_pack(
                product_dir,
                {"Record ID": "ASV-1", "Product Name": "Foo", "Shop Name": "Bar", "Product Link": "https://x"},
                "Shopee",
                transport=fake_transport,
            )
            self.assertIsNone(path)
            self.assertIsNotNone(err)


class HookOverlayTests(unittest.TestCase):
    def test_sanitize_hook_strips_quotes_and_trailing_period(self):
        from process_shopee_video import sanitize_hook

        self.assertEqual(sanitize_hook('  "Wait til you see this."  '), "Wait til you see this")
        self.assertEqual(sanitize_hook("'Stop scrolling now!'"), "Stop scrolling now!")
        # Curly quotes
        self.assertEqual(sanitize_hook("“Hook line”"), "Hook line")

    def test_sanitize_hook_keeps_hook_punctuation(self):
        from process_shopee_video import sanitize_hook

        self.assertEqual(sanitize_hook("You won't believe step 3?"), "You won't believe step 3?")

    def test_sanitize_hook_takes_first_line_only(self):
        from process_shopee_video import sanitize_hook

        self.assertEqual(sanitize_hook("Best deal\nThis is reasoning"), "Best deal")

    def test_generate_hook_returns_none_on_llm_failure(self):
        from process_shopee_video import generate_hook

        def fake_transport(url, payload, timeout):
            return 500, {"error": {"message": "down"}}

        result = generate_hook(
            {"Product Name": "Foo", "Shop Name": "Bar"}, "Shopee", transport=fake_transport
        )
        self.assertIsNone(result)

    def test_generate_hook_returns_sanitized_text(self):
        from process_shopee_video import generate_hook

        def fake_transport(url, payload, timeout):
            return 200, {"choices": [{"message": {"content": '"Stop scrolling now"'}}]}

        result = generate_hook(
            {"Product Name": "Foo", "Shop Name": "Bar"}, "Shopee", transport=fake_transport
        )
        self.assertEqual(result, "Stop scrolling now")

    def test_build_ffmpeg_command_with_hook_inserts_top_drawtext(self):
        # Use a video extension so we hit the video branch
        command = build_ffmpeg_command(
            "ffmpeg",
            Path("input.mp4"),
            Path("output.mp4"),
            "Shop - Product",
            render_mode="passthrough",
            hook_text="Wait for it",
        )
        graph = command[command.index("-filter_complex") + 1]
        # Hook drawtext is added on top of the existing one
        self.assertEqual(graph.count("drawtext="), 2)
        self.assertIn("Wait for it", graph)
        self.assertIn("between(t,0,1.5)", graph)
        self.assertIn("alpha=", graph)
        self.assertTrue(graph.endswith("[v]"))

    def test_build_ffmpeg_command_without_hook_unchanged(self):
        with_hook = build_ffmpeg_command(
            "ffmpeg", Path("input.mp4"), Path("output.mp4"), "cap", render_mode="passthrough"
        )
        graph = with_hook[with_hook.index("-filter_complex") + 1]
        self.assertEqual(graph.count("drawtext="), 1)
        self.assertNotIn("between(t,", graph)

    def test_required_columns_includes_hook(self):
        from process_shopee_video import REQUIRED_COLUMNS

        self.assertIn("Hook", REQUIRED_COLUMNS)


class CoverImageTests(unittest.TestCase):
    def test_cover_extract_time_uses_2s_for_long_video(self):
        from process_shopee_video import cover_extract_time

        self.assertEqual(cover_extract_time(60.0), 2.0)
        self.assertEqual(cover_extract_time(2.5), 2.0)

    def test_cover_extract_time_falls_back_for_short_video(self):
        from process_shopee_video import cover_extract_time

        self.assertEqual(cover_extract_time(1.0), 0.5)
        self.assertAlmostEqual(cover_extract_time(0.4), 0.2)
        # Unknown duration → safe default
        self.assertEqual(cover_extract_time(None), 0.5)

    def test_build_cover_command_includes_seek_and_drawtext_when_hook_set(self):
        from process_shopee_video import build_cover_ffmpeg_command

        cmd = build_cover_ffmpeg_command(
            "ffmpeg", Path("video.mp4"), Path("cover.jpg"), "Stop scrolling now", 2.0
        )
        self.assertIn("-ss", cmd)
        self.assertIn("2.00", cmd)
        self.assertIn("-frames:v", cmd)
        vf = cmd[cmd.index("-vf") + 1]
        self.assertIn("scale=1080:1920", vf)
        self.assertIn("drawtext=", vf)
        self.assertIn("Stop scrolling now", vf)

    def test_build_cover_command_skips_drawtext_when_no_hook(self):
        from process_shopee_video import build_cover_ffmpeg_command

        cmd = build_cover_ffmpeg_command("ffmpeg", Path("v.mp4"), Path("c.jpg"), None, 1.0)
        vf = cmd[cmd.index("-vf") + 1]
        self.assertNotIn("drawtext", vf)


class TelegramHelpersTests(unittest.TestCase):
    def test_parse_captions_sections_round_trip(self):
        sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
        from telegram_bot import parse_captions_sections

        content = (
            "# header\n# more headers\n\n"
            "=== TikTok ===\n"
            "Hook line + CTA #shopee\n\n"
            "=== Reels (Instagram) ===\n"
            "Different copy here\n\n"
            "=== Hashtags ===\n"
            "#a #b #c\n"
        )
        sections = parse_captions_sections(content)
        labels = [s[0] for s in sections]
        self.assertEqual(labels, ["TikTok", "Reels (Instagram)", "Hashtags"])
        self.assertEqual(sections[0][1], "Hook line + CTA #shopee")
        self.assertEqual(sections[2][1], "#a #b #c")

    def test_escape_md_v2_escapes_specials(self):
        sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
        from telegram_bot import escape_md_v2

        self.assertEqual(escape_md_v2("a.b!c"), "a\\.b\\!c")
        self.assertEqual(escape_md_v2("(test)"), "\\(test\\)")

    def test_escape_md_v2_code_only_escapes_backtick_and_backslash(self):
        sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
        from telegram_bot import escape_md_v2_code

        # Inside a code block, only ` and \ need escaping; periods and parens are fine.
        self.assertEqual(escape_md_v2_code("a.b!c"), "a.b!c")
        self.assertEqual(escape_md_v2_code("here is `code`"), "here is \\`code\\`")


if __name__ == "__main__":
    unittest.main()
